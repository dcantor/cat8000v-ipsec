#!/usr/bin/env python3
"""Build docs/requirements.xlsx from docs/requirements.html — the same source the PDF is rendered from, so the two never
diverge. Sheets: Overview (document metadata, the system context table, a count per area), Requirements (every numbered
requirement: id, area, section, priority, text, acceptance criteria — filterable), Roles & glossary, Traceability,
Constraints. Usage: build_xlsx.py [requirements.html]"""
import html as htmlmod
import re
import sys
from html.parser import HTMLParser
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
SRC = HERE / (sys.argv[1] if len(sys.argv) > 1 else "requirements.html")
OUT = SRC.with_suffix(".xlsx")

AREAS = {"NET": "Network service", "SOT": "Source of truth (Nautobot)", "INT": "Intent document", "NAC": "Network-as-Code pipeline",
         "FW": "Firewalls", "PKI": "IKE authentication / PKI", "CUST": "Customers and design patterns", "HOST": "LAN hosts",
         "INET": "Internet breakout", "SEC": "Access control and audit", "RUN": "Runs and day-2 operations", "UI": "Portal pages",
         "CMP": "Compliance and remediation", "API": "REST API", "MON": "Monitoring and alerting", "TST": "Verification (tests)",
         "LAB": "Lab infrastructure", "DOC": "Documentation", "NFR": "Non-functional",
         "DCI": "Interconnect, NAT and DNS"}


class Doc(HTMLParser):
    """The document as a flat list: ('h1'|'h2'|'table'|'list', payload). A table is a list of rows of cell texts; the first
    row is the header when the source marked it with <th>."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.items, self.text, self.table, self.row, self.cell, self.li = [], [], None, None, None, None
        self.depth_td = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("h1", "h2", "h3"): self.text = []
        elif tag == "table": self.table = []
        elif tag == "tr": self.row = []
        elif tag in ("td", "th"): self.cell = []; self.depth_td += 1
        elif tag == "ul" and self.table is None: self.li = []
        elif tag == "li" and self.li is not None: self.cell = []
        elif tag == "br" and self.cell is not None: self.cell.append(" ")

    def handle_endtag(self, tag):
        if tag in ("h1", "h2", "h3"): self.items.append((tag, clean("".join(self.text)))); self.text = []
        elif tag == "table":
            if self.table: self.items.append(("table", self.table))
            self.table = None
        elif tag == "tr":
            if self.row: self.table.append(self.row)
            self.row = None
        elif tag in ("td", "th"):
            if self.row is not None: self.row.append(clean("".join(self.cell or [])))
            self.cell = None; self.depth_td -= 1
        elif tag == "li" and self.li is not None and self.cell is not None:
            self.li.append(clean("".join(self.cell))); self.cell = None
        elif tag == "ul" and self.li is not None:
            if self.li: self.items.append(("list", self.li))
            self.li = None

    def handle_data(self, data):
        if self.cell is not None: self.cell.append(data)
        elif self.text is not None and self.table is None: self.text.append(data)


def clean(s):
    return re.sub(r"\s+", " ", htmlmod.unescape(s)).strip()


doc = Doc(); doc.feed(SRC.read_text())
title = next((t for k, t in doc.items if k == "h1"), "Requirements")
subtitle = ""

# ---- walk the document: the current h2 is the section every table belongs to --------------------------------------------
section = ""; meta, context, roles, glossary, traceability, constraints, reqs = [], [], [], [], [], [], []
for kind, payload in doc.items:
    if kind == "h2": section = payload; continue
    if kind == "list" and "Constraints" in section: constraints += payload; continue
    if kind != "table": continue
    head = [c.lower() for c in payload[0]]
    if not section:                                        # the metadata table under the title
        meta += [r for r in payload if len(r) == 2]
    elif head[:2] == ["id", "prio"]:                       # a requirement table
        for r in payload[1:]:
            if len(r) < 4: continue
            rid = r[0]; pre = rid.split("-")[0]
            reqs.append({"id": rid, "area": AREAS.get(pre, pre), "section": section, "prio": r[1], "req": r[2], "acc": r[3]})
    elif section.startswith("1.") and head[0] == "component": context = payload
    elif section.startswith("2.") and head[0] == "role": roles = payload
    elif section.startswith("2.") and head[0] == "term": glossary = payload
    elif "Traceability" in section: traceability = payload

# ---- workbook -------------------------------------------------------------------------------------------------------------
HEAD = PatternFill("solid", fgColor="EEF2F7"); BLUE = Font(bold=True, color="0B62D6")
THIN = Side(style="thin", color="D7DDE5"); BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOP = Alignment(vertical="top", wrap_text=True); TOPL = Alignment(vertical="top", wrap_text=False)
PRIO = {"MUST": PatternFill("solid", fgColor="FDE8E8"), "SHOULD": PatternFill("solid", fgColor="FFF3D6"), "MAY": PatternFill("solid", fgColor="E6F2FF")}
wb = Workbook(); wb.remove(wb.active)


def sheet(name, header, rows, widths, title_rows=()):
    ws = wb.create_sheet(name)
    r = 1
    for t, f in title_rows:
        ws.cell(r, 1, t).font = f; r += 1
    if title_rows: r += 1
    head_row = r
    for i, h in enumerate(header, 1):
        c = ws.cell(r, i, h); c.font = Font(bold=True); c.fill = HEAD; c.border = BOX; c.alignment = TOP
    for row in rows:
        r += 1
        for i, v in enumerate(row, 1):
            c = ws.cell(r, i, v); c.border = BOX; c.alignment = TOP
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(head_row + 1, 1)
    if rows: ws.auto_filter.ref = f"A{head_row}:{get_column_letter(len(header))}{r}"
    return ws, head_row


# Overview
ws = wb.create_sheet("Overview")
ws["A1"] = title; ws["A1"].font = Font(bold=True, size=18)
ws["A2"] = clean(re.search(r'<div class="sub">(.*?)</div>', SRC.read_text(), re.S).group(1)); ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
ws.merge_cells("A2:D2"); ws.row_dimensions[2].height = 32
r = 4
for k, v in meta:
    ws.cell(r, 1, k).font = Font(bold=True); ws.cell(r, 1).alignment = TOPL
    c = ws.cell(r, 2, v); c.alignment = TOP; ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4); ws.row_dimensions[r].height = 30; r += 1
r += 1
ws.cell(r, 1, "Requirements per area").font = BLUE; r += 1
for i, h in enumerate(["Area", "Requirements", "MUST", "SHOULD", "MAY"], 1):
    c = ws.cell(r, i, h); c.font = Font(bold=True); c.fill = HEAD; c.border = BOX
counts = {}
for q in reqs: counts.setdefault(q["area"], []).append(q["prio"])
for area, ps in counts.items():
    r += 1
    for i, v in enumerate([area, len(ps), ps.count("MUST"), ps.count("SHOULD"), ps.count("MAY")], 1):
        c = ws.cell(r, i, v); c.border = BOX; c.alignment = TOPL
r += 1
for i, v in enumerate(["Total", len(reqs), sum(1 for q in reqs if q["prio"] == "MUST"), sum(1 for q in reqs if q["prio"] == "SHOULD"), sum(1 for q in reqs if q["prio"] == "MAY")], 1):
    c = ws.cell(r, i, v); c.border = BOX; c.font = Font(bold=True)
r += 2
ws.cell(r, 1, "System context").font = BLUE; r += 1
for j, row in enumerate(context):
    for i, v in enumerate(row, 1):
        c = ws.cell(r, i, v); c.border = BOX; c.alignment = TOP
        if j == 0: c.font = Font(bold=True); c.fill = HEAD
    ws.row_dimensions[r].height = 28 if j else 15; r += 1
for col, w in zip("ABCDE", (30, 46, 46, 34, 12)): ws.column_dimensions[col].width = w
ws.freeze_panes = "A4"

# Requirements
ws, head = sheet("Requirements", ["ID", "Area", "Section", "Priority", "Requirement", "Acceptance criteria / verified by"],
                 [[q["id"], q["area"], q["section"], q["prio"], q["req"], q["acc"]] for q in reqs],
                 (11, 26, 34, 10, 78, 54))
for row in ws.iter_rows(min_row=head + 1, max_row=head + len(reqs)):
    row[0].font = BLUE
    row[3].fill = PRIO.get(row[3].value, PatternFill()); row[3].alignment = Alignment(vertical="top", horizontal="center")
    ws.row_dimensions[row[0].row].height = 58

# Roles & glossary
rows = [["Role", r0[0], r0[1]] for r0 in roles[1:]] + [["Term", g[0], g[1]] for g in glossary[1:]]
ws, head = sheet("Roles & glossary", ["Kind", "Name", "Meaning / needs"], rows, (10, 26, 110))
for row in ws.iter_rows(min_row=head + 1, max_row=head + len(rows)): ws.row_dimensions[row[0].row].height = 42

# Traceability
ws, head = sheet("Traceability", traceability[0] if traceability else ["Area", "Implemented in", "Verified by"], traceability[1:], (28, 76, 36))
for row in ws.iter_rows(min_row=head + 1, max_row=head + max(0, len(traceability) - 1)): ws.row_dimensions[row[0].row].height = 34

# Constraints
ws, head = sheet("Constraints", ["#", "Constraint, assumption or known limit"], [[i, t] for i, t in enumerate(constraints, 1)], (5, 140))
for row in ws.iter_rows(min_row=head + 1, max_row=head + len(constraints)): ws.row_dimensions[row[0].row].height = 60

wb.save(OUT)
print(f"{SRC.name} -> {OUT.name}: {len(reqs)} requirements in {len(counts)} areas, {len(roles) + len(glossary) - 2} glossary/role rows, "
      f"{max(0, len(traceability) - 1)} traceability rows, {len(constraints)} constraints ({OUT.stat().st_size / 1024:.0f} KB)")
